"""Panel verisinden tek slaytlik PPTX dashboard uretir.

Gorsel dil qkg-haftalik-rapor ile ayni: lacivert baslik, acik gri zemin,
beyaz kartlar (tile), native/duzenlenebilir grafikler.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY   = RGBColor(0x12, 0x1F, 0x3D)
PAGEBG = RGBColor(0xEE, 0xF1, 0xF6)
TILE   = RGBColor(0xFF, 0xFF, 0xFF)
TILELN = RGBColor(0xE4, 0xE9, 0xF0)
GREEN  = RGBColor(0x2E, 0x7D, 0x32)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREY   = RGBColor(0x6B, 0x76, 0x83)
DARKTX = RGBColor(0x22, 0x2B, 0x34)
BLUE   = RGBColor(0x1F, 0x6F, 0xD6)

HEAD_FONT = "Georgia"
BODY_FONT = "Calibri"


def _text(slide, x, y, w, h, s, size, color, *, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BODY_FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tb


def _tile(slide, x, y, w, h):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = 0.06
    sp.fill.solid()
    sp.fill.fore_color.rgb = TILE
    sp.line.color.rgb = TILELN
    sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def _style_chart(chart, colors):
    chart.font.size = Pt(9)
    chart.font.name = BODY_FONT
    chart.font.color.rgb = DARKTX
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    for series, color in zip(chart.plots[0].series, colors):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color


def _fmt_tr(iso: str) -> str:
    d = datetime.fromisoformat(iso)
    months = ["Oca", "Şub", "Mar", "Nis", "May", "Haz", "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]
    return f"{d.day} {months[d.month - 1]}"


def build_pptx(data: dict) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PAGEBG

    single_day = data["from"] == data["to"]
    range_text = _fmt_tr(data["from"]) if single_day else f'{_fmt_tr(data["from"])} – {_fmt_tr(data["to"])}'

    _text(slide, 0.55, 0.32, 9.5, 0.55, "Xray Test Koşum Raporu", 26, NAVY, bold=True, font=HEAD_FONT)
    _text(slide, 0.55, 0.92, 12.0, 0.3,
          f'{data["plan"]} · {range_text} · {data["executions"]} execution',
          12, GREY)
    _text(slide, 10.6, 0.45, 2.2, 0.3, datetime.now().strftime("%d.%m.%Y %H:%M"),
          10, GREY, align=PP_ALIGN.RIGHT)

    # ---- KPI kartlari ----
    total = data["total"]
    pr = data.get("progress") or {}
    pct = (lambda n: f"%{round(n / total * 100)}" if total else "–")
    prog_val = f'%{round(pr["done"] / pr["total"] * 100)}' if pr.get("total") else "–"
    prog_note = f'{pr.get("done", 0)}/{pr.get("total", 0)} test'
    if pr.get("estimateDays"):
        prog_note += f' · ~{pr["estimateDays"]} gün'
    kpis = [
        ("TOPLAM KOŞUM", str(total), range_text, DARKTX),
        ("BAŞARILI", str(data["passCount"]), pct(data["passCount"]), GREEN),
        ("BAŞARISIZ", str(data["failCount"]), pct(data["failCount"]), RED),
        ("KİŞİ", str(len(data["people"])), "test koşan", DARKTX),
        ("PLAN İLERLEMESİ", prog_val, prog_note, BLUE),
    ]
    kw, gap, kx, ky, kh = 2.25, 0.24, 0.55, 1.38, 1.12
    for i, (label, value, note, vcolor) in enumerate(kpis):
        x = kx + i * (kw + gap)
        _tile(slide, x, ky, kw, kh)
        _text(slide, x + 0.18, ky + 0.13, kw - 0.36, 0.25, label, 9, GREY, bold=True)
        _text(slide, x + 0.18, ky + 0.36, kw - 0.36, 0.45, value, 24, vcolor, bold=True)
        _text(slide, x + 0.18, ky + 0.83, kw - 0.36, 0.22, note, 8.5, GREY)

    # ---- Grafik kartlari ----
    cy, ch = 2.75, 4.25
    left_x, left_w = 0.55, 6.55
    right_x, right_w = 7.34, 5.44
    people = data["people"][:10]

    _tile(slide, left_x, cy, left_w, ch)
    _text(slide, left_x + 0.22, cy + 0.16, left_w - 0.44, 0.3, "Kişi Bazında Koşum", 13, DARKTX, bold=True)
    if people:
        cd = CategoryChartData()
        def short(name):
            base = name.split(" - ")[0].split(" (")[0].title().strip()
            return base[:30] + "…" if len(base) > 31 else base
        cd.categories = [short(p["name"]) for p in reversed(people)]
        cd.add_series("Başarılı", [p["pass"] for p in reversed(people)])
        cd.add_series("Başarısız", [p["fail"] for p in reversed(people)])
        cd.add_series("Diğer", [p["other"] for p in reversed(people)])
        gf = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_STACKED,
            Inches(left_x + 0.15), Inches(cy + 0.5), Inches(left_w - 0.3), Inches(ch - 0.68), cd)
        _style_chart(gf.chart, [GREEN, RED, GREY])
    else:
        _text(slide, left_x + 0.22, cy + 1.8, left_w - 0.44, 0.4, "Bu aralıkta koşum yok", 12, GREY,
              align=PP_ALIGN.CENTER)

    _tile(slide, right_x, cy, right_w, ch)
    if single_day:
        _text(slide, right_x + 0.22, cy + 0.16, right_w - 0.44, 0.3, "Sonuç Dağılımı", 13, DARKTX, bold=True)
        if total:
            cd = CategoryChartData()
            cats, vals, colors = [], [], []
            for label, key, color in (("Başarılı", "passCount", GREEN),
                                      ("Başarısız", "failCount", RED),
                                      ("Diğer", "otherCount", GREY)):
                if data[key]:
                    cats.append(label)
                    vals.append(data[key])
                    colors.append(color)
            cd.categories = cats
            cd.add_series("Koşum", vals)
            gf = slide.shapes.add_chart(
                XL_CHART_TYPE.PIE,
                Inches(right_x + 0.15), Inches(cy + 0.5), Inches(right_w - 0.3), Inches(ch - 0.68), cd)
            chart = gf.chart
            chart.font.size = Pt(9)
            chart.font.name = BODY_FONT
            chart.font.color.rgb = DARKTX
            chart.has_title = False
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            for pt, color in zip(chart.plots[0].series[0].points, colors):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = color
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.show_value = True
            plot.data_labels.font.size = Pt(9)
    else:
        _text(slide, right_x + 0.22, cy + 0.16, right_w - 0.44, 0.3, "Günlere Göre Koşum", 13, DARKTX, bold=True)
        days = data["days"]
        cd = CategoryChartData()
        cd.categories = [_fmt_tr(d["date"]) for d in days]
        cd.add_series("Başarılı", [d["pass"] for d in days])
        cd.add_series("Başarısız", [d["fail"] for d in days])
        cd.add_series("Diğer", [d["other"] for d in days])
        gf = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_STACKED,
            Inches(right_x + 0.15), Inches(cy + 0.5), Inches(right_w - 0.3), Inches(ch - 0.68), cd)
        _style_chart(gf.chart, [GREEN, RED, GREY])

    _text(slide, 0.55, 7.12, 12.2, 0.25,
          "Kaynak: jira.thy.com · Xray (Raven) REST API · yalnızca seçilen aralıkta tamamlanmış koşumlar sayılmıştır",
          8, GREY)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
